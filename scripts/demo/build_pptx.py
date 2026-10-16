import json, re, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
import pymupdf

D = sys.argv[1]; OUT = sys.argv[2]
geo = json.load(open(f"{D}/geo.json"))
E = 6350  # EMU per CSS px on a 1920-wide 16:9 slide
prs = Presentation(); prs.slide_width = Emu(1920 * E); prs.slide_height = Emu(1080 * E)
blank = prs.slide_layouts[6]

def rgb(css):
    m = re.findall(r"[\d.]+", css or "0,0,0")
    return RGBColor(int(float(m[0])), int(float(m[1])), int(float(m[2])))

for sd in geo:
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(f"{D}/bg/{sd['bg']}", 0, 0, prs.slide_width, prs.slide_height)
    for el in sd["els"]:
        x, y, w, h = (Emu(int(el[k] * E)) for k in ("x", "y", "w", "h"))
        if el["kind"] == "box":
            sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
            sh.adjustments[0] = min(0.5, 18 / min(el["w"], el["h"]))
            sh.fill.solid(); sh.fill.fore_color.rgb = rgb(el["fill"]); sh.line.fill.background(); sh.shadow.inherit = False
        elif el["kind"] == "img":
            path = f"{D}/{el['src']}"
            px = pymupdf.Pixmap(path); iw, ih = px.width, px.height
            if el["fit"] == "contain":
                s = min(el["w"] / iw, el["h"] / ih); dw, dh = iw * s, ih * s
                slide.shapes.add_picture(path, Emu(int((el["x"] + (el["w"] - dw) / 2) * E)), Emu(int((el["y"] + (el["h"] - dh) / 2) * E)), Emu(int(dw * E)), Emu(int(dh * E)))
            else:  # cover, anchored top
                pic = slide.shapes.add_picture(path, x, y, w, h)
                box_r, img_r = el["w"] / el["h"], iw / ih
                if img_r > box_r:
                    cut = (1 - box_r / img_r) / 2; pic.crop_left = cut; pic.crop_right = cut
                else:
                    pic.crop_bottom = 1 - img_r / box_r
        else:
            tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
            tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            first = True
            for para in el["paras"]:
                p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
                p.line_spacing = 1.1
                for run in para:
                    r = p.add_run(); r.text = run["t"]
                    f = r.font; f.name = "Arial"; f.size = Pt(run["size"] * 0.5); f.bold = bool(run["b"]); f.color.rgb = rgb(run["color"])
prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
